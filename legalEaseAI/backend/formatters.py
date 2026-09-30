from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt
from fpdf import FPDF


DISCLAIMER = (
    "AI-generated legal draft. "
    "Please consult a qualified lawyer before signing."
)


def sanitize_text(text):
    """Clean document text."""

    if not isinstance(text, str):
        raise TypeError("Text must be a string")

    return (
        text.replace("\r\n", "\n")
        .replace("\r", "\n")
        .strip()
    )


# -----------------------------
# TXT FORMATTER
# -----------------------------

def format_txt(text):

    clean = sanitize_text(text)

    content = clean + "\n\n" + DISCLAIMER

    return content.encode("utf-8")


# -----------------------------
# DOCX FORMATTER
# -----------------------------

def format_docx(text, doc_type="Legal Document"):

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    style = document.styles["Normal"]

    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    heading = document.add_heading(
        doc_type.upper(),
        level=0
    )

    heading.alignment = 1

    clean = sanitize_text(text)

    for line in clean.split("\n"):

        line = line.strip()

        if not line:
            continue

        if (
            line.isupper()
            or (
                len(line) < 80
                and line[:1].isdigit()
                and "." in line[:5]
            )
        ):

            document.add_heading(
                line,
                level=2
            )

        else:

            paragraph = document.add_paragraph(line)

            paragraph.paragraph_format.space_after = Pt(8)

    footer = section.footer.paragraphs[0]

    footer.text = DISCLAIMER

    footer.alignment = 1

    output = BytesIO()

    document.save(output)

    return output.getvalue()


# -----------------------------
# PDF FORMATTER
# -----------------------------

def format_pdf(text, doc_type="Legal Document"):

    pdf = FPDF()

    pdf.set_auto_page_break(
        auto=True,
        margin=20
    )

    pdf.add_page()

    pdf.set_margins(20, 20, 20)

    # Document title

    pdf.set_font("Times", "B", 16)

    pdf.multi_cell(
        0,
        10,
        doc_type.upper(),
        align="C"
    )

    pdf.ln(5)

    clean = sanitize_text(text)

    pdf.set_font("Times", size=12)

    for line in clean.split("\n"):

        line = line.strip()

        if not line:
            pdf.ln(4)
            continue

        safe_line = line.encode(
            "latin-1",
            "replace"
        ).decode("latin-1")

        if (
            safe_line.isupper()
            or (
                len(safe_line) < 80
                and safe_line[:1].isdigit()
                and "." in safe_line[:5]
            )
        ):

            pdf.set_font("Times", "B", 12)

            pdf.multi_cell(
                0,
                8,
                safe_line
            )

            pdf.set_font("Times", size=12)

        else:

            pdf.multi_cell(
                0,
                7,
                safe_line
            )

    # Footer

    pdf.set_y(-18)

    pdf.set_font("Times", "I", 8)

    pdf.cell(
        0,
        10,
        DISCLAIMER,
        align="C"
    )

    return bytes(pdf.output())